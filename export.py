# export.py
import os
import logging

log = logging.getLogger('export')

def get_desktop_path():
    for name in ("Desktop", "Рабочий стол"):
        p = os.path.join(os.path.expanduser("~"), name)
        if os.path.exists(p):
            return p
    return os.path.expanduser("~")

def _import_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise RuntimeError(
            "Библиотека openpyxl не установлена.\n"
            "Установите её командой: pip install openpyxl"
        )

def export_to_excel(db, filepath=None, only_rp_id=None, only_org_id=None):
    """Экспорт всей базы или выбранной ветки."""
    openpyxl = _import_openpyxl()
    from openpyxl.styles import Border, Side, Alignment, Font

    if not filepath:
        export_dir = os.path.join(get_desktop_path(), "Экспорт контактов")
        os.makedirs(export_dir, exist_ok=True)
        filepath = os.path.join(export_dir, "contacts_export.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Контакты"

    thin = Border(*[Side(style='thin')] * 4)
    bold = Font(bold=True)
    al = Alignment(horizontal='left', vertical='center')
    ac = Alignment(horizontal='center', vertical='center')

    ws.append(["Тип", "Наименование / ФИО", "Город / Договор",
               "Дата договора", "Адрес / Должность", "Телефон"])

    for rp in db.get_responsible_persons():
        if only_rp_id is not None and rp['id'] != only_rp_id and only_org_id is None:
            continue
        ws.append(["Ответственный", rp['full_name'], rp['city'] or '-', "",
                   rp['position'] or '-', rp['phone'] or '-'])
        for org in db.get_organizations(rp['id']):
            if only_org_id is not None and org['id'] != only_org_id:
                continue
            ws.append(["Организация", org['name'], org['contract_number'] or '-',
                       org['contract_date'] or '-', org['address'] or '-',
                       org['phone'] or '-'])
            for store in db.get_stores(org['id']):
                ws.append(["Магазин", store['name'], "", "",
                           store['unloading_address'] or '-', store['phone'] or '-'])
                for c in db.get_contacts(store['id'], 'decision_maker'):
                    ws.append(["ЛПР", c['full_name'] or '-', "", "",
                               c['position'] or '-', c['phone'] or '-'])
                for c in db.get_contacts(store['id'], 'payment'):
                    ws.append(["Оплата", c['full_name'] or '-', "", "",
                               c['position'] or '-', c['phone'] or '-'])

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.border = thin
            # Ошибка была здесь: cell.font = ... else cell.font
            if cell.row == 1:
                cell.font = bold
                cell.alignment = ac
            else:
                cell.alignment = al

    for col, w in zip('ABCDEF', (20, 35, 20, 15, 35, 20)):
        ws.column_dimensions[col].width = w

    wb.save(filepath)
    log.info(f"Exported to {filepath}")
    return filepath

def import_from_excel(db, filepath, skip_duplicates=True):
    """Импорт из Excel. Возвращает (добавлено, пропущено_ошибок, пропущено_дубликатов)."""
    openpyxl = _import_openpyxl()
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    added = skipped = duplicates = 0
    cur_rp = cur_org = cur_store = None

    def find_rp(name):
        db.cursor.execute("SELECT id FROM responsible_persons WHERE full_name=? AND deleted_at IS NULL", (name,))
        res = db.cursor.fetchone()
        return res['id'] if res else None

    def find_org(rp_id, name):
        db.cursor.execute("SELECT id FROM organizations WHERE rp_id=? AND name=? AND deleted_at IS NULL", (rp_id, name))
        res = db.cursor.fetchone()
        return res['id'] if res else None

    def find_store(org_id, name):
        db.cursor.execute("SELECT id FROM stores WHERE org_id=? AND name=? AND deleted_at IS NULL", (org_id, name))
        res = db.cursor.fetchone()
        return res['id'] if res else None

    def find_contact(store_id, ctype, full_name, pos, phone):
        db.cursor.execute(
            "SELECT id FROM contacts WHERE store_id=? AND contact_type=? AND "
            "COALESCE(full_name,'')=COALESCE(?,'') AND "
            "COALESCE(position,'')=COALESCE(?,'') AND "
            "COALESCE(phone,'')=COALESCE(?,'') AND deleted_at IS NULL",
            (store_id, ctype, full_name, pos, phone))
        return db.cursor.fetchone() is not None

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            skipped += 1
            continue

        typ = str(row[0] or '').strip()
        # Оборачиваем в str(), чтобы числа (телефоны, даты) не ломали .strip()
        name = str(row[1] or '').strip() if len(row) > 1 else ''
        c2   = str(row[2] or '').strip() if len(row) > 2 else ''
        c3   = str(row[3] or '').strip() if len(row) > 3 else ''
        c4   = str(row[4] or '').strip() if len(row) > 4 else ''
        c5   = str(row[5] or '').strip() if len(row) > 5 else ''

        try:
            if typ == 'Ответственный':
                existing = find_rp(name)
                if existing and skip_duplicates:
                    cur_rp = existing
                    duplicates += 1
                else:
                    cur_rp = db.add_responsible_person(name, c2, c4, c5)
                    added += 1
                cur_org = cur_store = None

            elif typ == 'Организация':
                if cur_rp is None:
                    skipped += 1
                    continue
                existing = find_org(cur_rp, name)
                if existing and skip_duplicates:
                    cur_org = existing
                    duplicates += 1
                else:
                    cur_org = db.add_organization(cur_rp, name, c4, c5, c2, c3)
                    added += 1
                cur_store = None

            elif typ == 'Магазин':
                if cur_org is None:
                    skipped += 1
                    continue
                existing = find_store(cur_org, name)
                if existing and skip_duplicates:
                    cur_store = existing
                    duplicates += 1
                else:
                    cur_store = db.add_store(cur_org, name, c4, c5)
                    added += 1

            elif typ in ('ЛПР', 'Оплата'):
                if cur_store is None:
                    skipped += 1
                    continue
                ctype = 'decision_maker' if typ == 'ЛПР' else 'payment'
                if skip_duplicates and find_contact(cur_store, ctype, name, c4, c5):
                    duplicates += 1
                else:
                    db.add_contact(cur_store, ctype, name, c4, c5)
                    added += 1
            else:
                skipped += 1
        except Exception as e:
            log.warning(f"Import row skipped: {e}")
            skipped += 1

    log.info(f"Imported: added={added}, skipped={skipped}, duplicates={duplicates}")
    return added, skipped, duplicates